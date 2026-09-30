from fastapi import APIRouter, HTTPException, Header
from fastapi.responses import StreamingResponse
from app.services.firebase_service import verify_token, db
from app.services.assistant_service import calculate_win_probability
from app.services.memory_service import get_user_memory
from firebase_admin import firestore
import io

router = APIRouter()

def get_uid(authorization: str):
    try:
        token   = authorization.replace("Bearer ", "")
        decoded = verify_token(token)
        return decoded["uid"]
    except Exception as e:
        raise HTTPException(status_code=401, detail=str(e))


@router.get("/")
def get_funnel(authorization: str = Header(...)):
    """
    Returns enriched funnel data — every lead with
    meeting count, task count, last meeting summary,
    follow-up status and win probability.
    """
    uid    = get_uid(authorization)
    memory = get_user_memory(uid)

    leads    = memory["leads"]
    meetings = memory["meetings"]
    tasks    = memory["tasks"]
    followups= memory["follow_ups"]

    funnel = []
    for lead in leads:
        lead_id      = lead.get("id","")
        company      = (lead.get("company") or "").lower()

        # Find linked meetings
        linked_meetings = [
            m for m in meetings
            if m.get("linked_lead_id") == lead_id or
               company in (m.get("linked_company") or "").lower()
        ]
        linked_meetings.sort(
            key=lambda x: x.get("created_at",""), reverse=True
        )

        # Find linked tasks
        linked_tasks = [
            t for t in tasks
            if t.get("lead_id") == lead_id
        ]
        pending_tasks = [
            t for t in linked_tasks
            if t.get("status") != "done"
        ]

        # Find follow-ups
        linked_followups = [
            f for f in followups
            if (f.get("company") or "").lower() == company
        ]
        overdue_followups = [
            f for f in linked_followups
            if f.get("status") == "pending"
        ]

        # Last meeting data
        last_meeting     = linked_meetings[0] if linked_meetings else None
        last_meeting_date= last_meeting.get("created_at","")[:10] if last_meeting else ""
        last_meeting_summary = (last_meeting.get("summary","")[:120] + "...") \
            if last_meeting and last_meeting.get("summary") else ""

        # Win probability
        win_prob = calculate_win_probability(lead, memory)

        funnel.append({
            "id":                   lead_id,
            "company":              lead.get("company",""),
            "contact_name":         lead.get("contact_name",""),
            "email":                lead.get("email",""),
            "phone":                lead.get("phone",""),
            "stage":                lead.get("stage","lead"),
            "value":                lead.get("value") or 0,
            "win_probability":      win_prob,
            "last_contact":         lead.get("last_contact",""),
            "notes":                lead.get("notes",""),
            "meetings_count":       len(linked_meetings),
            "last_meeting_date":    last_meeting_date,
            "last_meeting_summary": last_meeting_summary,
            "total_tasks":          len(linked_tasks),
            "pending_tasks":        len(pending_tasks),
            "follow_ups":           len(linked_followups),
            "overdue_followups":    len(overdue_followups),
            "created_at":           lead.get("created_at",""),
            "source":               lead.get("source","manual"),
        })

    # Sort by stage order then value
    stage_order = [
        "negotiation","proposal","qualified",
        "prospect","lead","closed_won","closed_lost"
    ]
    funnel.sort(key=lambda x: (
        stage_order.index(x["stage"])
        if x["stage"] in stage_order else 99,
        -(x["value"] or 0)
    ))

    return funnel


@router.get("/export")
def export_funnel_excel(authorization: str = Header(...)):
    """Export funnel as formatted Excel file."""
    uid    = get_uid(authorization)
    memory = get_user_memory(uid)
    leads  = memory["leads"]
    meetings = memory["meetings"]
    tasks    = memory["tasks"]

    import openpyxl
    from openpyxl.styles import (
        Font, PatternFill, Alignment,
        Border, Side, GradientFill
    )
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales Funnel"

    # Colors
    BLUE       = "0078D4"
    DARK_BLUE  = "1B3A6B"
    LIGHT_BLUE = "EFF6FF"
    GREEN      = "16A34A"
    AMBER      = "D97706"
    RED        = "DC2626"
    GRAY       = "F5F5F5"
    WHITE      = "FFFFFF"

    STAGE_COLORS = {
        "lead":        "EEF2FF",
        "prospect":    "EFF6FF",
        "qualified":   "ECFEFF",
        "proposal":    "FFF7ED",
        "negotiation": "FEF2F2",
        "closed_won":  "F0FDF4",
        "closed_lost": "F9FAFB",
    }

    thin = Side(style="thin", color="E5E5E5")
    border = Border(top=thin, left=thin, right=thin, bottom=thin)

    # ── Title row ──
    ws.merge_cells("A1:R1")
    title_cell = ws["A1"]
    title_cell.value = "SALES FUNNEL SHEET — AI Sales Copilot"
    title_cell.font  = Font(
        name="Calibri", bold=True, size=14, color=WHITE
    )
    title_cell.fill  = PatternFill(
        "solid", fgColor=DARK_BLUE
    )
    title_cell.alignment = Alignment(
        horizontal="center", vertical="center"
    )
    ws.row_dimensions[1].height = 32

    # ── Headers ──
    headers = [
        "Company", "Contact", "Email", "Phone",
        "Stage", "Deal Value ($)", "Win %",
        "Last Contact", "Meetings", "Last Meeting Date",
        "Last Meeting Summary", "Total Tasks",
        "Pending Tasks", "Follow-ups", "Overdue F/U",
        "Notes", "Source", "Added"
    ]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col, value=header)
        cell.font      = Font(name="Calibri", bold=True,
                              size=10, color=WHITE)
        cell.fill      = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(
            horizontal="center", vertical="center",
            wrap_text=True
        )
        cell.border    = border
    ws.row_dimensions[2].height = 24

    # ── Data rows ──
    stage_order = [
        "negotiation","proposal","qualified",
        "prospect","lead","closed_won","closed_lost"
    ]

    all_leads = []
    for lead in leads:
        lead_id  = lead.get("id","")
        company  = (lead.get("company") or "").lower()
        linked_m = [
            m for m in meetings
            if m.get("linked_lead_id") == lead_id or
               company in (m.get("linked_company") or "").lower()
        ]
        linked_m.sort(
            key=lambda x: x.get("created_at",""), reverse=True
        )
        linked_t = [t for t in tasks if t.get("lead_id") == lead_id]

        win_prob = calculate_win_probability(lead, memory)
        last_m   = linked_m[0] if linked_m else None

        all_leads.append({
            **lead,
            "win_probability":      win_prob,
            "meetings_count":       len(linked_m),
            "last_meeting_date":    last_m.get("created_at","")[:10] if last_m else "",
            "last_meeting_summary": (last_m.get("summary","")[:100] + "...") if last_m else "",
            "total_tasks":          len(linked_t),
            "pending_tasks":        len([t for t in linked_t if t.get("status")!="done"]),
        })

    all_leads.sort(key=lambda x: (
        stage_order.index(x.get("stage","lead"))
        if x.get("stage","lead") in stage_order else 99,
        -(x.get("value") or 0)
    ))

    for row_num, lead in enumerate(all_leads, 3):
        stage     = lead.get("stage","lead")
        row_color = STAGE_COLORS.get(stage, WHITE)
        win_prob  = lead.get("win_probability", 0)

        row_data = [
            lead.get("company",""),
            lead.get("contact_name",""),
            lead.get("email","") or "",
            lead.get("phone","") or "",
            stage.replace("_"," ").title(),
            lead.get("value") or 0,
            f"{win_prob}%",
            lead.get("last_contact","") or "",
            lead.get("meetings_count", 0),
            lead.get("last_meeting_date","") or "",
            lead.get("last_meeting_summary","") or "",
            lead.get("total_tasks", 0),
            lead.get("pending_tasks", 0),
            lead.get("follow_ups", 0) if "follow_ups" in lead else 0,
            lead.get("overdue_followups", 0) if "overdue_followups" in lead else 0,
            (lead.get("notes","") or "")[:200],
            lead.get("source","manual"),
            lead.get("created_at","")[:10] if lead.get("created_at") else "",
        ]

        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col, value=value)
            cell.font      = Font(name="Calibri", size=9)
            cell.fill      = PatternFill("solid", fgColor=row_color)
            cell.border    = border
            cell.alignment = Alignment(
                vertical="center", wrap_text=True
            )

            # Win probability color
            if col == 7:
                color = GREEN if win_prob >= 70 \
                    else AMBER if win_prob >= 40 else RED
                cell.font = Font(
                    name="Calibri", size=9,
                    bold=True, color=color
                )

            # Deal value formatting
            if col == 6 and value:
                cell.number_format = '#,##0'
                cell.alignment     = Alignment(horizontal="right")

        ws.row_dimensions[row_num].height = 32

    # ── Column widths ──
    col_widths = [
        22, 18, 28, 14, 14, 14, 8,
        14, 10, 16, 45, 10, 12, 10, 10,
        40, 12, 12
    ]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # ── Freeze header rows ──
    ws.freeze_panes = "A3"

    # ── Summary sheet ──
    ws2 = wb.create_sheet("Pipeline Summary")
    ws2["A1"] = "Stage"
    ws2["B1"] = "Count"
    ws2["C1"] = "Total Value ($)"
    ws2["D1"] = "Avg Win %"

    for cell in ["A1","B1","C1","D1"]:
        ws2[cell].font  = Font(bold=True, color=WHITE)
        ws2[cell].fill  = PatternFill("solid", fgColor=BLUE)
        ws2[cell].alignment = Alignment(horizontal="center")

    stage_summary = {}
    for lead in all_leads:
        s = lead.get("stage","lead")
        if s not in stage_summary:
            stage_summary[s] = {"count":0,"value":0,"prob_sum":0}
        stage_summary[s]["count"]    += 1
        stage_summary[s]["value"]    += lead.get("value") or 0
        stage_summary[s]["prob_sum"] += lead.get("win_probability",0)

    for row, (stage, data) in enumerate(stage_summary.items(), 2):
        avg_prob = round(data["prob_sum"] / data["count"]) \
            if data["count"] else 0
        ws2.cell(row=row, column=1, value=stage.replace("_"," ").title())
        ws2.cell(row=row, column=2, value=data["count"])
        ws2.cell(row=row, column=3, value=data["value"])\
           .number_format = '#,##0'
        ws2.cell(row=row, column=4, value=f"{avg_prob}%")
        color = STAGE_COLORS.get(stage, WHITE)
        for col in range(1,5):
            ws2.cell(row=row, column=col).fill = \
                PatternFill("solid", fgColor=color)

    for col, width in zip(["A","B","C","D"], [20,10,18,12]):
        ws2.column_dimensions[col].width = width

    # ── Save ──
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": 'attachment; filename="Sales_Funnel.xlsx"'
        }
    )